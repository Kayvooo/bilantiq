"""Tabellen aus CSV- und Excel-Dateien lesen (Bankimport, Stammdaten-Import).

CSV-Dateien unterscheiden sich in Kodierung, Trennzeichen, Datums- und Zahlenformat. Deshalb
wird hier nur erkannt und vorgeschlagen, die Spaltenbedeutung legt der Aufrufer bzw. der
Nutzer fest. Excel-Dateien (.xlsx) werden mit openpyxl gelesen: nur gespeicherte Werte,
Formeln werden nie ausgewertet, und das Archiv wird vor dem Oeffnen auf Zip-Bomben geprueft."""
from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from typing import Literal

MAX_ZEILEN = 20000
DATUMSFORMATE = ("%d.%m.%Y", "%Y-%m-%d", "%d.%m.%y", "%d/%m/%Y", "%m/%d/%Y", "%Y/%m/%d", "%d-%m-%Y")
TRENNZEICHEN = (";", ",", "\t", "|")

# Grenzen fuer Excel-Archive (entpackt), damit eine kleine Datei nicht Speicher und CPU flutet.
XLSX_MAX_EINTRAEGE = 1000
XLSX_MAX_ENTPACKT = 50 * 1024 * 1024
XLSX_MAX_KOMPRESSION = 200

_OLE_SIGNATUR = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


class TabellenFehler(ValueError):
    """Die Datei kann nicht gelesen werden. Die Meldung ist fuer den Nutzer gedacht."""


@dataclass
class Tabelle:
    kopf: list[str]
    zeilen: list[list[str]]
    format: Literal["csv", "xlsx"]
    kodierung: str | None = None
    trennzeichen: str | None = None
    blatt: str | None = None
    blaetter: list[str] = field(default_factory=list)


# ---------------------------------------------------------------------------
# CSV


def dekodiere(daten: bytes, kodierung: str | None = None) -> tuple[str, str]:
    if kodierung:
        try:
            return daten.decode(kodierung), kodierung
        except (UnicodeDecodeError, LookupError) as exc:
            raise TabellenFehler(f"Die Datei lässt sich nicht als {kodierung} lesen.") from exc
    for versuch in ("utf-8-sig", "cp1252"):
        try:
            return daten.decode(versuch), versuch
        except UnicodeDecodeError:
            continue
    return daten.decode("latin-1"), "latin-1"


def erkenne_trennzeichen(text: str) -> str:
    zeilen = [zeile for zeile in text.splitlines() if zeile.strip()][:10]
    bestes, beste_spalten = ";", 1
    for trenner in TRENNZEICHEN:
        spalten = [len(zeile) for zeile in csv.reader(zeilen, delimiter=trenner)]
        if spalten and len(set(spalten)) == 1 and spalten[0] > beste_spalten:
            bestes, beste_spalten = trenner, spalten[0]
    return bestes


def _eindeutige_koepfe(kopf: list[str]) -> list[str]:
    eindeutig: list[str] = []
    for nummer, name in enumerate(kopf, start=1):
        name = name or f"Spalte {nummer}"
        kandidat, zaehler = name, 2
        while kandidat in eindeutig:
            kandidat, zaehler = f"{name} ({zaehler})", zaehler + 1
        eindeutig.append(kandidat)
    return eindeutig


def zeilen_aus_text(
    text: str, trennzeichen: str, *, leer_meldung: str = "Die Datei enthält keine Datenzeilen.",
) -> tuple[list[str], list[list[str]]]:
    leser = csv.reader(io.StringIO(text), delimiter=trennzeichen)
    zeilen = [[feld.strip() for feld in zeile] for zeile in leser if any(feld.strip() for feld in zeile)]
    if len(zeilen) < 2:
        raise TabellenFehler(leer_meldung)
    if len(zeilen) - 1 > MAX_ZEILEN:
        raise TabellenFehler(f"Die Datei enthält mehr als {MAX_ZEILEN} Zeilen. Bitte in kleinere Teile aufteilen.")
    return _eindeutige_koepfe(zeilen[0]), zeilen[1:]


# ---------------------------------------------------------------------------
# Werte


def parse_datum(wert: str, format_: str) -> date | None:
    wert = wert.strip()
    for versuch in (wert, wert.split()[0] if wert else ""):
        try:
            return datetime.strptime(versuch, format_).date()
        except ValueError:
            continue
    return None


def parse_betrag(wert: str, dezimal: str) -> Decimal | None:
    text = wert.strip().replace(" ", "").replace("\xa0", "")
    negativ = text.startswith("(") and text.endswith(")")
    text = re.sub(r"[^\d,.\-+]", "", text)
    if not text:
        return None
    if dezimal == "komma":
        text = text.replace(".", "").replace(",", ".")
    else:
        text = text.replace(",", "")
    try:
        zahl = Decimal(text)
    except InvalidOperation:
        return None
    return -abs(zahl) if negativ else zahl


def schlage_datumsformat_vor(werte: list[str]) -> str:
    echte = [wert for wert in werte if wert.strip()]
    for format_ in DATUMSFORMATE:
        if echte and all(parse_datum(wert, format_) is not None for wert in echte):
            return format_
    return DATUMSFORMATE[0]


def schlage_dezimal_vor(werte: list[str]) -> Literal["komma", "punkt"]:
    komma = sum(1 for wert in werte if re.search(r",\d{1,2}\s*[^\d]*$", wert.strip()))
    punkt = sum(1 for wert in werte if re.search(r"\.\d{1,2}\s*[^\d]*$", wert.strip()))
    return "punkt" if punkt > komma else "komma"


# ---------------------------------------------------------------------------
# Excel


def _pruefe_xlsx_archiv(daten: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(daten)) as archiv:
            eintraege = archiv.infolist()
    except zipfile.BadZipFile as exc:
        raise TabellenFehler("Die Excel-Datei ist beschädigt.") from exc
    if len(eintraege) > XLSX_MAX_EINTRAEGE:
        raise TabellenFehler("Die Excel-Datei enthält zu viele Teile.")
    if sum(eintrag.file_size for eintrag in eintraege) > XLSX_MAX_ENTPACKT:
        raise TabellenFehler("Die Excel-Datei ist entpackt zu groß. Bitte als CSV speichern oder aufteilen.")
    for eintrag in eintraege:
        if eintrag.file_size > 1024 * 1024 and eintrag.file_size > XLSX_MAX_KOMPRESSION * max(eintrag.compress_size, 1):
            raise TabellenFehler("Die Excel-Datei ist ungewöhnlich stark gepackt und wird nicht geöffnet.")
    if "xl/workbook.xml" not in {eintrag.filename for eintrag in eintraege}:
        raise TabellenFehler("Das ist keine Excel-Datei (.xlsx). Bitte in Excel oder LibreOffice als .xlsx oder .csv speichern.")


def _zahl_als_text(zahl: Decimal) -> str:
    # Excel-Zahlen werden im deutschen Format ausgegeben (Komma, ohne Tausenderpunkt), damit
    # sie zu Textzellen im deutschen Format passen.
    text = format(zahl.normalize(), "f") if zahl == zahl.to_integral_value() else format(zahl, "f")
    return text.replace(".", ",")


def _zelle_als_text(wert: object) -> str:
    if wert is None:
        return ""
    if isinstance(wert, bool):
        return "ja" if wert else "nein"
    if isinstance(wert, datetime):
        return wert.strftime("%d.%m.%Y") if wert.time() == time(0, 0) else wert.strftime("%d.%m.%Y %H:%M")
    if isinstance(wert, date):
        return wert.strftime("%d.%m.%Y")
    if isinstance(wert, int):
        return str(wert)
    if isinstance(wert, float):
        return _zahl_als_text(Decimal(repr(wert)))
    if isinstance(wert, Decimal):
        return _zahl_als_text(wert)
    return str(wert).strip()


def _lese_xlsx(daten: bytes, blatt: str | None) -> Tabelle:
    _pruefe_xlsx_archiv(daten)
    from openpyxl import load_workbook

    try:
        mappe = load_workbook(io.BytesIO(daten), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl wirft je nach Defekt sehr verschiedene Fehler
        raise TabellenFehler("Die Excel-Datei lässt sich nicht lesen. Bitte als .xlsx oder .csv neu speichern.") from exc
    try:
        namen = list(mappe.sheetnames)
        if not namen:
            raise TabellenFehler("Die Excel-Datei enthält kein Tabellenblatt.")
        name = blatt if blatt in namen else namen[0]
        roh: list[list[str]] = []
        for zeile in mappe[name].iter_rows(values_only=True):
            werte = [_zelle_als_text(wert) for wert in zeile]
            if not any(werte):
                continue
            roh.append(werte)
            if len(roh) > MAX_ZEILEN + 1:
                raise TabellenFehler(f"Das Blatt enthält mehr als {MAX_ZEILEN} Zeilen. Bitte in kleinere Teile aufteilen.")
    finally:
        mappe.close()

    if len(roh) < 2:
        raise TabellenFehler(f"Das Blatt „{name}“ enthält keine Datenzeilen.")
    kopf = roh[0]
    while kopf and not kopf[-1]:
        kopf = kopf[:-1]
    breite = len(kopf)
    zeilen = [(zeile + [""] * breite)[:breite] for zeile in roh[1:]]
    return Tabelle(
        kopf=_eindeutige_koepfe(kopf), zeilen=[z for z in zeilen if any(z)], format="xlsx", blatt=name, blaetter=namen,
    )


# ---------------------------------------------------------------------------
# Einstieg


def lese_tabelle(
    daten: bytes, *, blatt: str | None = None, kodierung: str | None = None, trennzeichen: str | None = None,
) -> Tabelle:
    """Erkennt CSV oder Excel am Inhalt, nicht an Dateiname oder Content-Type."""
    if daten.startswith(_OLE_SIGNATUR):
        raise TabellenFehler(
            "Alte Excel-Dateien (.xls) oder passwortgeschützte Mappen werden nicht unterstützt. "
            "Bitte in Excel als .xlsx oder .csv speichern."
        )
    if daten.startswith(b"PK\x03\x04"):
        return _lese_xlsx(daten, blatt)
    if daten.startswith(b"%PDF") or b"\x00" in daten[:4096]:
        raise TabellenFehler("Bitte eine CSV- oder Excel-Datei hochladen.")
    text, erkannte_kodierung = dekodiere(daten, kodierung)
    trenner = trennzeichen or erkenne_trennzeichen(text)
    kopf, zeilen = zeilen_aus_text(text, trenner)
    breite = len(kopf)
    zeilen = [(zeile + [""] * breite)[:breite] if len(zeile) != breite else zeile for zeile in zeilen]
    return Tabelle(kopf=kopf, zeilen=zeilen, format="csv", kodierung=erkannte_kodierung, trennzeichen=trenner)
