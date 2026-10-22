import io
import json

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.db import SessionLocal
from app.invoices.kontakte import kontakt_iban_hash
from app.invoices.models import Kontakt
from tests.test_benutzerverwaltung import _register_and_verify
from tests.test_invoices import _invoice_payload, _onboarded_mandant

KONTAKTE_CSV = (
    "Kundennr.;Firma;Vorname;Nachname;Straße;Hausnummer;PLZ;Ort;Land;E-Mail;Telefon;USt-IdNr.;IBAN;Zahlungsziel;Faxnummer\n"
    "K-1;Müller GmbH;Erika;Muster;Hauptstraße;5;1067;Dresden;Deutschland;info@mueller.de;+49 351 1234;DE123456789;"
    "DE89 3704 0044 0532 0130 00;30 Tage;0351 999\n"
    "K-2;;Max;Mustermann;Weg;1;80331;München;DE;max@example.com;;;;;\n"
    "K-3;Ohne Mail AG;;;;;10115;Berlin;Atlantis;kaputt@;;;DE00 1234;;\n"
    "K-4;;;;;;;;;;;;;;\n"
    "K-1;Müller GmbH;;;;;;;;;;;;;\n"
).encode("cp1252")

ZUORDNUNG = {
    "kundennummer": "Kundennr.", "name": "Firma", "vorname": "Vorname", "nachname": "Nachname", "strasse": "Straße",
    "hausnummer": "Hausnummer", "plz": "PLZ", "ort": "Ort", "land": "Land", "email": "E-Mail", "telefon": "Telefon",
    "ust_id": "USt-IdNr.", "iban": "IBAN", "zahlungsziel_tage": "Zahlungsziel",
}


def _post(client, headers, pfad, daten: bytes, optionen: dict | None = None, dateiname="kontakte.csv"):
    form = {"optionen": json.dumps(optionen)} if optionen is not None else {}
    return client.post(pfad, headers=headers, files={"file": (dateiname, daten, "text/csv")}, data=form)


def _optionen(**extra) -> dict:
    return {"zuordnung": ZUORDNUNG, **extra}


def test_vorschau_schlaegt_zuordnung_vor(client: TestClient):
    headers = _onboarded_mandant(client)
    resp = _post(client, headers, "/stammdaten/kontakte/import/vorschau", KONTAKTE_CSV)
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["anzahl_zeilen"] == 5
    assert body["kodierung"] == "cp1252"
    assert body["vorgeschlagene_zuordnung"] == ZUORDNUNG
    assert any(feld["schluessel"] == "name" and feld["pflicht"] for feld in body["felder"])


def test_probelauf_zeigt_status_je_zeile_und_legt_nichts_an(client: TestClient):
    headers = _onboarded_mandant(client)
    resp = _post(client, headers, "/stammdaten/kontakte/import/pruefen", KONTAKTE_CSV, _optionen())
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert (body["neu"], body["fehler"], body["doppelt"], body["vorhanden"]) == (3, 1, 1, 0)
    zeilen = {z["nummer"]: z for z in body["zeilen"]}
    assert zeilen[5]["status"] == "fehler" and "Name fehlt." in zeilen[5]["fehler"]
    assert zeilen[6]["status"] == "doppelt" and "Zeile 2" in zeilen[6]["grund"]
    # Optionale Felder mit Fehlern bleiben leer, die Zeile wird trotzdem angelegt.
    hinweise = " ".join(zeilen[4]["hinweise"])
    assert "Atlantis" in hinweise and "E-Mail" in hinweise and "IBAN" in hinweise
    assert client.get("/kontakte", headers=headers).json() == []

    beispiel = body["beispiele"][0]
    assert beispiel["Name"] == "Müller GmbH"
    assert beispiel["PLZ"] == "01067"
    assert beispiel["Straße"] == "Hauptstraße 5"
    assert beispiel["Zahlungsziel (Tage)"] == "30"


def test_fehler_blockieren_bis_ausdruecklich_ausgelassen(client: TestClient):
    headers = _onboarded_mandant(client)
    gesperrt = _post(client, headers, "/stammdaten/kontakte/import", KONTAKTE_CSV, _optionen())
    assert gesperrt.status_code == 400
    assert "1 Zeilen haben Fehler" in gesperrt.json()["detail"]
    assert client.get("/kontakte", headers=headers).json() == []

    resp = _post(client, headers, "/stammdaten/kontakte/import", KONTAKTE_CSV, _optionen(fehlerhafte_auslassen=True))
    assert resp.status_code == 200, resp.text
    assert resp.json()["neu"] == 3

    kontakte = {k["name"]: k for k in client.get("/kontakte", headers=headers).json()}
    assert set(kontakte) == {"Müller GmbH", "Max Mustermann", "Ohne Mail AG"}
    mueller = kontakte["Müller GmbH"]
    assert (mueller["vorname"], mueller["nachname"], mueller["plz"], mueller["land"]) == ("Erika", "Muster", "01067", "DE")
    assert mueller["iban"] == "DE89370400440532013000"
    assert kontakte["Max Mustermann"]["ist_geschaeftskunde"] is False
    assert kontakte["Ohne Mail AG"]["email"] is None and kontakte["Ohne Mail AG"]["land"] == "DE"
    with SessionLocal() as db:
        zeile = db.get(Kontakt, mueller["id"])
        assert zeile.iban_hash == kontakt_iban_hash(zeile.mandant_id, "DE89370400440532013000")
        assert zeile.import_id == resp.json()["import_id"]


def test_reimport_erkennt_bestand_und_ergaenzen_ueberschreibt_nie(client: TestClient):
    headers = _onboarded_mandant(client)
    client.post("/kontakte", headers=headers, json={"name": "Müller GmbH", "kundennummer": "K-1", "telefon": "0000", "ort": "Alt"})
    datei = (
        "Kundennr.;Firma;Ort;Telefon;E-Mail\n"
        "K-1;Müller GmbH;Dresden;+49 351 1234;info@mueller.de\n"
        "K-9;Neu GmbH;Leipzig;;\n"
    ).encode("utf-8")
    zuordnung = {"kundennummer": "Kundennr.", "name": "Firma", "ort": "Ort", "telefon": "Telefon", "email": "E-Mail"}

    ueberspringen = _post(client, headers, "/stammdaten/kontakte/import/pruefen", datei, {"zuordnung": zuordnung}).json()
    assert (ueberspringen["neu"], ueberspringen["vorhanden"]) == (1, 1)

    ergaenzen = _post(client, headers, "/stammdaten/kontakte/import", datei, {"zuordnung": zuordnung, "doppelte": "ergaenzen"})
    assert ergaenzen.status_code == 200, ergaenzen.text
    zeile = next(z for z in ergaenzen.json()["zeilen"] if z["status"] == "ergaenzen")
    assert zeile["ergaenzt"] == ["E-Mail"]

    mueller = next(k for k in client.get("/kontakte", headers=headers).json() if k["name"] == "Müller GmbH")
    assert (mueller["ort"], mueller["telefon"], mueller["email"]) == ("Alt", "0000", "info@mueller.de")


def test_gleiche_kundennummer_bei_anderem_namen_ist_ein_fehler(client: TestClient):
    headers = _onboarded_mandant(client)
    client.post("/kontakte", headers=headers, json={"name": "Ganz Andere KG", "kundennummer": "K-1"})
    datei = "Kundennr.;Firma\nK-1;Müller GmbH\n".encode()
    body = _post(client, headers, "/stammdaten/kontakte/import/pruefen", datei, {"zuordnung": {"kundennummer": "Kundennr.", "name": "Firma"}}).json()
    assert body["fehler"] == 1
    assert "gehört bereits zu „Ganz Andere KG“" in body["zeilen"][0]["fehler"][0]


def test_rueckgaengig_loescht_nur_unbenutzte(client: TestClient):
    headers = _onboarded_mandant(client)
    datei = "Firma;PLZ;Ort;Straße\nErste GmbH;10115;Berlin;Weg 1\nZweite GmbH;10115;Berlin;Weg 2\n".encode()
    zuordnung = {"name": "Firma", "plz": "PLZ", "ort": "Ort", "strasse": "Straße"}
    import_id = _post(client, headers, "/stammdaten/kontakte/import", datei, {"zuordnung": zuordnung}).json()["import_id"]
    erste = next(k for k in client.get("/kontakte", headers=headers).json() if k["name"] == "Erste GmbH")
    client.post("/invoices", headers=headers, json=_invoice_payload(erste["id"]))

    resp = client.post(f"/stammdaten/importe/{import_id}/rueckgaengig", headers=headers)
    assert resp.status_code == 200
    assert resp.json() == {"geloescht": 1, "behalten": 1}
    assert [k["name"] for k in client.get("/kontakte", headers=headers).json()] == ["Erste GmbH"]
    assert client.post(f"/stammdaten/importe/{import_id}/rueckgaengig", headers=headers).status_code == 409

    protokoll = client.get("/stammdaten/importe", headers=headers).json()
    assert protokoll[0]["id"] == import_id and protokoll[0]["rueckgaengig_am"] is not None


def test_produkte_import_mit_bruttopreis_und_standard_ust(client: TestClient):
    headers = _onboarded_mandant(client)
    datei = _xlsx_datei([
        ["Art.-Nr.", "Artikelbezeichnung", "VK brutto", "MwSt %", "Einheit", "Einkaufspreis"],
        ["A-1", "Buch", 10.70, 7, "Stück", 5],
        ["A-2", "Beratung", 119, None, "Stunde", 0],
        ["A-3", "Ohne Preis", None, 19, None, None],
    ])
    vorschau = _post(client, headers, "/stammdaten/produkte/import/vorschau", datei, dateiname="artikel.xlsx").json()
    assert vorschau["vorgeschlagene_zuordnung"] == {
        "artikelnummer": "Art.-Nr.", "bezeichnung": "Artikelbezeichnung", "bruttopreis": "VK brutto", "ust_satz": "MwSt %",
        "einheit": "Einheit",
    }
    resp = _post(
        client, headers, "/stammdaten/produkte/import", datei,
        {"zuordnung": vorschau["vorgeschlagene_zuordnung"], "standard_ust_satz": "19"}, dateiname="artikel.xlsx",
    )
    assert resp.status_code == 200, resp.text
    produkte = {p["bezeichnung"]: p for p in client.get("/produkte", headers=headers).json()}
    assert (produkte["Buch"]["einzelpreis"], produkte["Buch"]["ust_satz"]) == ("10.00", "7.00")
    assert (produkte["Beratung"]["einzelpreis"], produkte["Beratung"]["ust_satz"]) == ("100.00", "19.00")
    assert produkte["Ohne Preis"]["einzelpreis"] == "0.00" and produkte["Ohne Preis"]["einheit"] == "Stück"
    assert produkte["Buch"]["artikelnummer"] == "A-1"


def _xlsx_datei(zeilen) -> bytes:
    mappe = Workbook()
    for zeile in zeilen:
        mappe.active.append(zeile)
    puffer = io.BytesIO()
    mappe.save(puffer)
    return puffer.getvalue()


def test_export_und_reimport_ergibt_nichts_neues(client: TestClient):
    headers = _onboarded_mandant(client)
    client.post("/kontakte", headers=headers, json={
        "name": "Formel GmbH", "telefon": "+49 30 1", "notiz": "=HYPERLINK(\"http://boese.example\")", "typ": "beide",
        "kundennummer": "K-7", "skonto_prozent": "2", "zahlungsziel_tage": 10,
    })
    client.post("/kontakte", headers=headers, json={"name": "Zweite AG", "plz": "10115"})

    for format_ in ("csv", "xlsx"):
        export = client.get(f"/stammdaten/kontakte/export?format={format_}", headers=headers)
        assert export.status_code == 200
        assert export.headers["cache-control"] == "no-store"
        vorschau = _post(client, headers, "/stammdaten/kontakte/import/vorschau", export.content, dateiname=f"x.{format_}").json()
        pruefung = _post(client, headers, "/stammdaten/kontakte/import/pruefen", export.content, {"zuordnung": vorschau["vorgeschlagene_zuordnung"]}).json()
        assert (pruefung["neu"], pruefung["vorhanden"], pruefung["fehler"]) == (0, 2, 0), (format_, pruefung)

    csv_text = client.get("/stammdaten/kontakte/export?format=csv", headers=headers).content.decode("utf-8-sig")
    # Die Notiz enthaelt Anfuehrungszeichen und steht deshalb selbst in Quotes.
    assert '"\'=HYPERLINK(' in csv_text and ";'+49 30 1;" in csv_text
    assert "Kunde und Lieferant" in csv_text and "2;" in csv_text

    xlsx = load_workbook(io.BytesIO(client.get("/stammdaten/kontakte/export?format=xlsx", headers=headers).content))
    zellen = [zelle for zeile in xlsx.active.iter_rows() for zelle in zeile]
    formel = next(z for z in zellen if isinstance(z.value, str) and "HYPERLINK" in z.value)
    assert formel.data_type == "s"


def test_vorlage_laesst_sich_direkt_pruefen(client: TestClient):
    headers = _onboarded_mandant(client)
    for ziel in ("kontakte", "produkte"):
        vorlage = client.get(f"/stammdaten/{ziel}/vorlage?format=csv", headers=headers)
        assert vorlage.status_code == 200
        vorschau = _post(client, headers, f"/stammdaten/{ziel}/import/vorschau", vorlage.content).json()
        pruefung = _post(client, headers, f"/stammdaten/{ziel}/import/pruefen", vorlage.content, {"zuordnung": vorschau["vorgeschlagene_zuordnung"]}).json()
        assert (pruefung["neu"], pruefung["fehler"], pruefung["mit_hinweisen"]) == (1, 0, 0), (ziel, pruefung)


def test_nur_inhaber_und_nur_eigener_mandant(client: TestClient):
    owner = _onboarded_mandant(client, "owner@example.com")
    code = client.post("/auth/einladungen", headers=owner, json={"rolle": "mitglied"}).json()["code"]
    mitglied = _register_and_verify(client, "mitglied@example.com", einladungscode=code)
    assert _post(client, mitglied, "/stammdaten/kontakte/import/vorschau", KONTAKTE_CSV).status_code == 403
    assert client.get("/stammdaten/kontakte/export", headers=mitglied).status_code == 403

    import_id = _post(client, owner, "/stammdaten/kontakte/import", KONTAKTE_CSV, _optionen(fehlerhafte_auslassen=True)).json()["import_id"]
    fremd = _onboarded_mandant(client, "fremd@example.com")
    assert client.post(f"/stammdaten/importe/{import_id}/rueckgaengig", headers=fremd).status_code == 404
    assert client.get("/stammdaten/importe", headers=fremd).json() == []
    assert "Müller" not in client.get("/stammdaten/kontakte/export", headers=fremd).content.decode("utf-8-sig")
    # Der Bestand des fremden Mandanten zaehlt nicht als "vorhanden".
    assert _post(client, fremd, "/stammdaten/kontakte/import/pruefen", KONTAKTE_CSV, _optionen()).json()["vorhanden"] == 0


def test_zuordnung_wird_geprueft(client: TestClient):
    headers = _onboarded_mandant(client)
    falsch = _post(client, headers, "/stammdaten/kontakte/import/pruefen", KONTAKTE_CSV, {"zuordnung": {"name": "Gibt es nicht"}})
    assert falsch.status_code == 400 and "Gibt es nicht" in falsch.json()["detail"]
    doppelt = _post(client, headers, "/stammdaten/kontakte/import/pruefen", KONTAKTE_CSV, {"zuordnung": {"name": "Firma", "ort": "Firma"}})
    assert "mehreren Feldern" in doppelt.json()["detail"]
    ohne_name = _post(client, headers, "/stammdaten/kontakte/import/pruefen", KONTAKTE_CSV, {"zuordnung": {"ort": "Ort"}})
    assert "Namen" in ohne_name.json()["detail"]


def test_fehlerliste_enthaelt_originalzeile_und_meldung(client: TestClient):
    headers = _onboarded_mandant(client)
    resp = _post(client, headers, "/stammdaten/kontakte/import/fehlerliste", KONTAKTE_CSV, _optionen())
    assert resp.status_code == 200
    text = resp.content.decode("utf-8-sig")
    zeilen = text.strip().split("\r\n")
    assert zeilen[0].endswith(";Fehler")
    assert any(zeile.startswith("K-4;") and zeile.endswith("Name fehlt.") for zeile in zeilen)
    assert any("Atlantis" in zeile for zeile in zeilen)
