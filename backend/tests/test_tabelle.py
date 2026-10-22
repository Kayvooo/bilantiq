import io
import zipfile
from datetime import date

import pytest
from openpyxl import Workbook

from app.core import tabelle as tabelle_modul
from app.core.tabelle import TabellenFehler, lese_tabelle


def _xlsx(zeilen: list[list[object]], blattname: str = "Kunden") -> bytes:
    mappe = Workbook()
    blatt = mappe.active
    blatt.title = blattname
    for zeile in zeilen:
        blatt.append(zeile)
    puffer = io.BytesIO()
    mappe.save(puffer)
    return puffer.getvalue()


def test_csv_windows_1252_mit_semikolon():
    daten = "Firma;Ort\nMüller GmbH;Köln\nBäckerei Groß;München\n".encode("cp1252")
    tabelle = lese_tabelle(daten)
    assert (tabelle.format, tabelle.kodierung, tabelle.trennzeichen) == ("csv", "cp1252", ";")
    assert tabelle.kopf == ["Firma", "Ort"]
    assert tabelle.zeilen == [["Müller GmbH", "Köln"], ["Bäckerei Groß", "München"]]


def test_csv_kurze_zeilen_werden_aufgefuellt_und_leere_koepfe_benannt():
    tabelle = lese_tabelle(b"Name;;Ort\nA;x;Berlin\nB\n")
    assert tabelle.kopf == ["Name", "Spalte 2", "Ort"]
    assert tabelle.zeilen[1] == ["B", "", ""]


def test_xlsx_werte_werden_zu_deutschem_text():
    daten = _xlsx([
        ["Firma", "PLZ", "Seit", "Preis", "Aktiv", None],
        ["Muster GmbH", 1067, date(2026, 10, 1), 90.5, True, None],
        [None, None, None, None, None, None],
        ["Zweite AG", "80331", None, 1200, False, None],
    ])
    tabelle = lese_tabelle(daten)
    assert tabelle.format == "xlsx"
    assert tabelle.blatt == "Kunden"
    assert tabelle.kopf == ["Firma", "PLZ", "Seit", "Preis", "Aktiv"]
    assert tabelle.zeilen == [
        ["Muster GmbH", "1067", "01.10.2026", "90,5", "ja"],
        ["Zweite AG", "80331", "", "1200", "nein"],
    ]


def test_xlsx_formeln_werden_nie_ausgewertet():
    daten = _xlsx([["Name", "Summe"], ["A", "=1+1"], ["B", "=HYPERLINK(\"http://example.com\")"]])
    tabelle = lese_tabelle(daten)
    # openpyxl speichert keinen berechneten Wert, data_only liefert daher nichts statt der Formel.
    assert [zeile[1] for zeile in tabelle.zeilen] == ["", ""]


def test_xlsx_zweites_blatt_waehlbar():
    mappe = Workbook()
    mappe.active.append(["Leer"])
    zweites = mappe.create_sheet("Produkte")
    zweites.append(["Bezeichnung"])
    zweites.append(["Beratung"])
    puffer = io.BytesIO()
    mappe.save(puffer)
    tabelle = lese_tabelle(puffer.getvalue(), blatt="Produkte")
    assert tabelle.blaetter == ["Sheet", "Produkte"]
    assert tabelle.zeilen == [["Beratung"]]


def test_altes_xls_und_pdf_werden_abgelehnt():
    with pytest.raises(TabellenFehler, match=r"\.xls"):
        lese_tabelle(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100)
    with pytest.raises(TabellenFehler, match="CSV- oder Excel"):
        lese_tabelle(b"%PDF-1.7\n...")


def _zip(eintraege: dict[str, bytes]) -> bytes:
    puffer = io.BytesIO()
    with zipfile.ZipFile(puffer, "w", compression=zipfile.ZIP_DEFLATED) as archiv:
        for name, inhalt in eintraege.items():
            archiv.writestr(name, inhalt)
    return puffer.getvalue()


def test_zip_ohne_arbeitsmappe_ist_keine_excel_datei():
    with pytest.raises(TabellenFehler, match="keine Excel-Datei"):
        lese_tabelle(_zip({"content.xml": b"<x/>", "mimetype": b"application/vnd.oasis.opendocument.spreadsheet"}))


def test_zip_bomben_werden_vor_dem_oeffnen_abgelehnt(monkeypatch):
    with pytest.raises(TabellenFehler, match="stark gepackt"):
        lese_tabelle(_zip({"xl/workbook.xml": b"<x/>", "xl/worksheets/sheet1.xml": b"0" * (3 * 1024 * 1024)}))
    with pytest.raises(TabellenFehler, match="zu viele Teile"):
        lese_tabelle(_zip({"xl/workbook.xml": b"<x/>", **{f"x/{n}.xml": b"" for n in range(1001)}}))
    monkeypatch.setattr(tabelle_modul, "XLSX_MAX_ENTPACKT", 1000)
    with pytest.raises(TabellenFehler, match="entpackt zu groß"):
        lese_tabelle(_zip({"xl/workbook.xml": b"<x/>" * 300}))


def test_kaputte_xlsx_gibt_verstaendliche_meldung():
    with pytest.raises(TabellenFehler, match="lässt sich nicht lesen"):
        lese_tabelle(_zip({"xl/workbook.xml": b"kein xml"}))
